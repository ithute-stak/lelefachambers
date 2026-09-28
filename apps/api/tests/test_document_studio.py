from decimal import Decimal
from io import BytesIO

from docx import Document

from app.document_studio import DemandLetterRequest, build_demand_letter, build_docx_bytes
from app.legal_recovery import CalculationMethod, DebtVerificationRequest, PaymentInput


def verified_request(**overrides):
    verification = DebtVerificationRequest(
        debtor_name="Mpho Moletsane",
        national_id="123456789",
        source_reference="BAT-001",
        loan_date="2026-01-01",
        principal=Decimal("1000"),
        annual_rate_percent=Decimal("20"),
        term_months=3,
        method=CalculationMethod.MICRO_LOAN,
        payments=[PaymentInput(amount=Decimal("464"), paid_on="2026-02-01")],
        claimed_outstanding=Decimal("928"),
        as_of_date="2026-02-01",
    )
    payload = dict(
        verification=verification,
        client_name="Batlokoa Financial Services",
        client_reference="CLIENT-77",
        chambers_reference="LC/BAT/2026/001",
        letter_date="2026-02-10",
        demand_days=7,
    )
    payload.update(overrides)
    return DemandLetterRequest(**payload)


def test_verified_debt_generates_demand_letter_html():
    result = build_demand_letter(verified_request())

    assert result.eligible is True
    assert result.deadline.isoformat() == "2026-02-17"
    assert result.amount_summary["verified_outstanding"] == "M 928.00"
    assert result.html is not None
    assert "LELEFA CHAMBERS" in result.html
    assert "MPHO MOLETSANE" in result.html
    assert result.signatory_name == "Advocate Mats'epe Lelefa, LLM"
    assert "Lenyora House, Office No. 4" in result.html


def test_approved_docx_template_preserves_letterhead_and_footer():
    payload, result = build_docx_bytes(verified_request())
    assert result.eligible is True
    document = Document(BytesIO(payload))

    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    header_text = "\n".join(paragraph.text for section in document.sections for paragraph in section.header.paragraphs)
    footer_text = "\n".join(paragraph.text for section in document.sections for paragraph in section.footer.paragraphs)

    assert "MPHO MOLETSANE" in body
    assert "National ID: 123456789" in body
    assert "Source Ref: BAT-001" in body
    assert "LC/BAT/2026/001" in body
    assert "M 928.00" in body + table_text
    assert "Lenyora House, Office No. 4" in header_text
    assert "LELEFA CHAMBERS" in footer_text
    assert "+266 5776 3829" in footer_text


def test_unverified_claim_never_renders_formal_letter():
    req = verified_request()
    req.verification.claimed_outstanding = Decimal("950")
    result = build_demand_letter(req)

    assert result.eligible is False
    assert result.html is None
    assert result.blockers


def test_fully_settled_account_is_blocked_from_document_generation():
    req = verified_request()
    req.verification.payments = [PaymentInput(amount=Decimal("1392"), paid_on="2026-02-01")]
    req.verification.claimed_outstanding = Decimal("0")
    result = build_demand_letter(req)

    assert result.eligible is False
    assert result.html is None
    assert any("fully settled" in item for item in result.blockers)
