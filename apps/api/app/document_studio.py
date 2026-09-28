from __future__ import annotations

import hashlib
import io
import os
import re
import secrets
import shutil
import subprocess
import tempfile
from datetime import date, timedelta
from html import escape
from pathlib import Path
from typing import Literal

from docx import Document
from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.legal_recovery import DebtVerificationRequest, DebtVerificationResponse, calculate, money
from app.main import User, audit, current_user, get_db
from app.recovery import MatterDocument, require_recovery, vault_path

router = APIRouter(prefix="/api/v1/document-studio", tags=["document-studio"])

TEMPLATE_PATH = Path(__file__).resolve().parent / "templates" / "batlokoa_formal_demand.docx"


class DemandLetterRequest(BaseModel):
    verification: DebtVerificationRequest
    client_name: str = Field(min_length=2, max_length=240)
    client_reference: str | None = Field(default=None, max_length=120)
    chambers_reference: str = Field(min_length=1, max_length=120)
    debtor_address: str | None = Field(default=None, max_length=600)
    demand_days: int = Field(default=7, ge=1, le=90)
    payment_instructions: str | None = Field(default=None, max_length=1200)
    additional_notice: str | None = Field(default=None, max_length=1200)
    signatory_name: str = Field(default="Advocate Mats'epe Lelefa, LLM", min_length=2, max_length=240)
    signatory_title: str = Field(default="Managing Partner", min_length=2, max_length=160)
    letter_date: date = Field(default_factory=date.today)


class SaveDemandLetterRequest(DemandLetterRequest):
    matter_id: int = Field(gt=0)
    visibility: Literal["internal", "client"] = "internal"
    formats: list[Literal["docx", "pdf"]] = Field(default_factory=lambda: ["docx", "pdf"])


class DemandLetterPreview(BaseModel):
    eligible: bool
    blockers: list[str]
    warnings: list[str]
    verification: DebtVerificationResponse
    subject: str
    salutation: str
    body_paragraphs: list[str]
    amount_summary: dict[str, str]
    deadline: date
    signatory_name: str
    signatory_title: str
    html: str | None = None


def _format_money(value) -> str:
    return f"M {money(value):,.2f}"


def _long_date(value: date) -> str:
    return value.strftime("%d %B %Y")


def _rate_description(req: DemandLetterRequest) -> str:
    rate = req.verification.annual_rate_percent
    rendered = f"{rate.normalize()}%" if rate != rate.to_integral() else f"{int(rate)}%"
    if req.verification.method.value == "micro_loan":
        return f"{rendered} contractual rate under the LoanHub Micro Loan Method"
    return f"{rendered} per annum (contractual)"


def _safe_stem(req: DemandLetterRequest) -> str:
    source = re.sub(r"[^A-Za-z0-9._-]+", "-", req.verification.source_reference).strip("-") or "DEBT"
    return f"{source}_{req.verification.debtor_name.replace(' ', '_')}_Formal_Demand"


def build_demand_letter(req: DemandLetterRequest) -> DemandLetterPreview:
    verification = calculate(req.verification)
    deadline = req.letter_date + timedelta(days=req.demand_days)
    snapshot = verification.snapshot
    outstanding = _format_money(snapshot.verified_outstanding)
    payments = _format_money(snapshot.payments_credited)
    principal = _format_money(snapshot.principal)

    subject = f"RE: FORMAL DEMAND FOR PAYMENT OF OUTSTANDING LOAN - {outstanding}"
    salutation = "Dear Sir/Madam,"
    rate_description = _rate_description(req)

    paragraphs = [
        f"(1) We act on the instructions of {req.client_name}, our Client, in relation to the loan account held in your name.",
        f"(2) Our Client's records reflect that, on or about {_long_date(req.verification.loan_date)}, a loan facility in the principal amount of {principal} was advanced to you.",
        f"(3) The facility attracted interest at the agreed {rate_description}. Payments totalling {payments} have been credited to the account. Following application of those payments and contractual interest accrued under the loan agreement, the independently verified amount presently due and owing is {outstanding}.",
        "(4) Despite prior recovery efforts undertaken on behalf of our Client, the above balance remains unpaid.",
        f"(5) We therefore demand that, within {req.demand_days} ({_number_word(req.demand_days)}) days of receipt of this letter, you either settle the outstanding balance of {outstanding} in full or contact Lelefa Chambers in writing to propose a payment arrangement acceptable to our Client.",
        "(6) If payment or a satisfactory written arrangement is not made within that period, we are instructed to advise our Client on further recovery action, which may include appropriate legal proceedings for recovery of the debt, together with contractual interest and recoverable legal costs, to the extent lawfully claimable.",
        f"(7) If you dispute the balance or contend that any payment has not been credited, please provide documentary proof to our Chambers within the same {req.demand_days}-day period. This letter constitutes a formal demand and is issued without waiver of any rights or remedies available to our Client.",
    ]

    if req.payment_instructions:
        paragraphs.append(f"Payment instructions: {req.payment_instructions.strip()}")
    if req.additional_notice:
        paragraphs.append(req.additional_notice.strip())

    amount_summary = {
        "principal": principal,
        "contractual_interest": _format_money(snapshot.contractual_interest),
        "processing_fee": _format_money(snapshot.processing_fee),
        "payments_credited": payments,
        "verified_outstanding": outstanding,
    }

    html = None
    if verification.demand_letter_eligible:
        address = "" if not req.debtor_address else f"<div class='debtor-address'>{escape(req.debtor_address)}</div>"
        client_ref = "" if not req.client_reference else f"<div><strong>Client Ref:</strong> {escape(req.client_reference)}</div>"
        paragraph_html = "".join(f"<p>{escape(p)}</p>" for p in paragraphs)
        html = f"""<!doctype html>
<html>
<head>
<meta charset='utf-8'>
<title>{escape(subject)}</title>
<style>
@page {{ size: A4; margin: 25mm 18mm 24mm; }}
* {{ box-sizing: border-box; }}
body {{ margin: 0; color: #111; font-family: Georgia, 'Times New Roman', serif; font-size: 14px; line-height: 1.38; }}
.page {{ min-height: 249mm; position: relative; padding-bottom: 22mm; }}
.letterhead {{ text-align: center; border-bottom: 1.5px solid #b78a34; padding: 0 0 5px; margin-bottom: 8px; }}
.letterhead img {{ width: 430px; max-width: 78%; height: auto; display: block; margin: 0 auto -3px; }}
.contact {{ color: #163b5c; font-family: Georgia, 'Times New Roman', serif; font-size: 10.5px; white-space: nowrap; }}
.meta-grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 2px 18px; margin: 5px 0 12px; }}
.meta-right {{ text-align: right; }}
.meta-grid strong {{ color: #173b5c; }}
.debtor-address {{ white-space: pre-line; margin-top: 3px; }}
.subject {{ color: #173b5c; font-weight: 700; font-size: 15px; margin: 12px 0 9px; }}
p {{ margin: 0 0 8px; text-align: justify; }}
.summary-title {{ color: #173b5c; font-size: 16px; font-weight: 700; margin: 9px 0 4px; }}
.summary {{ width: 88%; border-collapse: collapse; font-size: 12.5px; }}
.summary td {{ border: 1px solid #555; padding: 3px 7px; }}
.summary td.label {{ background: #173b5c; color: white; font-weight: 700; width: 21%; }}
.signature {{ margin-top: 17px; }}
.signature-line {{ width: 280px; border-top: 1px solid #333; margin: 38px 0 4px; }}
.signature strong {{ color: #173b5c; }}
.footer {{ position: absolute; bottom: 0; left: 10%; width: 80%; text-align: center; border-top: 1.5px solid #d4b66c; color: #173b5c; font-size: 7px; padding-top: 3px; white-space: nowrap; }}
@media print {{ .page {{ min-height: 249mm; }} }}
</style>
</head>
<body>
<div class='page'>
<header class='letterhead'>
<img src='/brand/lelefa-chambers-logo.svg' alt='Lelefa Chambers'>
<div class='contact'>Lenyora House, Office No. 4, 190 Nightingale Road, New Europa, Maseru, Lesotho&nbsp;&nbsp; | &nbsp;&nbsp;Matsepelelefa15@gmail.com&nbsp;&nbsp; | &nbsp;&nbsp;+266 5776 3829</div>
</header>
<section class='meta-grid'>
<div><strong>TO:</strong><br><strong>{escape(req.verification.debtor_name.upper())}</strong><br>National ID: {escape(req.verification.national_id)}<br>Source Ref: {escape(req.verification.source_reference)}{address}</div>
<div class='meta-right'><strong>{escape(_long_date(req.letter_date))}</strong><br><br><strong>Our Ref:</strong><br>{escape(req.chambers_reference)}{client_ref}</div>
</section>
<div class='subject'>{escape(subject)}</div>
{paragraph_html}
<div class='summary-title'>ACCOUNT SUMMARY</div>
<table class='summary'>
<tr><td class='label'>Loan date</td><td>{escape(_long_date(req.verification.loan_date))}</td><td class='label'>Original principal</td><td>{principal}</td></tr>
<tr><td class='label'>Interest rate</td><td>{escape(rate_description)}</td><td class='label'>Payments credited</td><td>{payments}</td></tr>
<tr><td class='label'>Outstanding</td><td>{outstanding}</td><td class='label'>Demand period</td><td>{req.demand_days} days from receipt</td></tr>
</table>
<div class='signature'>Yours faithfully,<div class='signature-line'></div><strong>{escape(req.signatory_name)}</strong></div>
<footer class='footer'>LELEFA CHAMBERS&nbsp;&nbsp; | &nbsp;&nbsp;Lenyora House, Office No. 4, 190 Nightingale Road, New Europa, Maseru&nbsp;&nbsp; | &nbsp;&nbsp;+266 5776 3829</footer>
</div>
</body>
</html>"""

    return DemandLetterPreview(
        eligible=verification.demand_letter_eligible,
        blockers=verification.blockers,
        warnings=verification.warnings,
        verification=verification,
        subject=subject,
        salutation=salutation,
        body_paragraphs=paragraphs,
        amount_summary=amount_summary,
        deadline=deadline,
        signatory_name=req.signatory_name,
        signatory_title=req.signatory_title,
        html=html,
    )


def _number_word(value: int) -> str:
    words = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six", 7: "seven", 8: "eight", 9: "nine", 10: "ten", 14: "fourteen", 21: "twenty-one", 30: "thirty"}
    return words.get(value, str(value))


def _replace_in_paragraph(paragraph, old: str, new: str) -> None:
    if not old or old == new:
        return
    while old in "".join(run.text for run in paragraph.runs):
        full = "".join(run.text for run in paragraph.runs)
        start = full.index(old)
        end = start + len(old)
        cursor = 0
        start_idx = end_idx = None
        start_offset = end_offset = 0
        for idx, run in enumerate(paragraph.runs):
            next_cursor = cursor + len(run.text)
            if start_idx is None and start < next_cursor:
                start_idx = idx
                start_offset = start - cursor
            if end <= next_cursor:
                end_idx = idx
                end_offset = end - cursor
                break
            cursor = next_cursor
        if start_idx is None or end_idx is None:
            return
        if start_idx == end_idx:
            text = paragraph.runs[start_idx].text
            paragraph.runs[start_idx].text = text[:start_offset] + new + text[end_offset:]
        else:
            first = paragraph.runs[start_idx]
            last = paragraph.runs[end_idx]
            first.text = first.text[:start_offset] + new
            for idx in range(start_idx + 1, end_idx):
                paragraph.runs[idx].text = ""
            last.text = last.text[end_offset:]


def _replace_everywhere(document: Document, old: str, new: str) -> None:
    for paragraph in document.paragraphs:
        _replace_in_paragraph(paragraph, old, new)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_paragraph(paragraph, old, new)


def _template_replacements(req: DemandLetterRequest, preview: DemandLetterPreview) -> list[tuple[str, str]]:
    snapshot = preview.verification.snapshot
    rate_description = _rate_description(req)
    outstanding = _format_money(snapshot.verified_outstanding)
    return [
        ("NTHAKO LIMAKATSO", req.verification.debtor_name.upper()),
        ("National ID: ____________________", f"National ID: {req.verification.national_id}"),
        ("Source Ref: B1406", f"Source Ref: {req.verification.source_reference}"),
        ("27 September 2026", _long_date(req.letter_date)),
        ("LC/BATL/B1406/270926", req.chambers_reference),
        ("Batlokoa, our Client", f"{req.client_name}, our Client"),
        ("05 July 2024", _long_date(req.verification.loan_date)),
        ("M 10,000.00", _format_money(snapshot.principal)),
        ("20% per annum", rate_description),
        ("M 1,600.00", _format_money(snapshot.payments_credited)),
        ("M __________", outstanding),
        ("M __________.", f"{outstanding}."),
        ("7 days from receipt", f"{req.demand_days} days from receipt"),
        ("seven (7) days", f"{_number_word(req.demand_days)} ({req.demand_days}) days"),
        ("seven (7)-day", f"{req.demand_days}-day"),
        ("Advocate Mats'epe Lelefa, LLM", req.signatory_name),
    ]


def build_docx_bytes(req: DemandLetterRequest) -> tuple[bytes, DemandLetterPreview]:
    preview = build_demand_letter(req)
    if not preview.eligible:
        raise HTTPException(status_code=409, detail={"message": "Demand letter is blocked by debt verification", "blockers": preview.blockers})
    if not TEMPLATE_PATH.exists():
        raise HTTPException(status_code=500, detail="Approved demand-letter template is unavailable")

    document = Document(str(TEMPLATE_PATH))
    for old, new in _template_replacements(req, preview):
        _replace_everywhere(document, old, new)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue(), preview


def build_pdf_bytes(req: DemandLetterRequest) -> tuple[bytes, DemandLetterPreview]:
    docx_bytes, preview = build_docx_bytes(req)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise HTTPException(status_code=503, detail="PDF renderer is unavailable on this server")
    with tempfile.TemporaryDirectory(prefix="lelefa-document-studio-") as temp_dir:
        workdir = Path(temp_dir)
        source = workdir / "demand.docx"
        source.write_bytes(docx_bytes)
        env = os.environ.copy()
        env["HOME"] = str(workdir / "home")
        Path(env["HOME"]).mkdir(parents=True, exist_ok=True)
        process = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(workdir), str(source)],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=60,
            env=env,
            check=False,
        )
        pdf = workdir / "demand.pdf"
        if process.returncode != 0 or not pdf.exists() or pdf.stat().st_size == 0:
            raise HTTPException(status_code=500, detail="Unable to render the approved demand-letter template to PDF")
        return pdf.read_bytes(), preview


def _store_generated(*, db: Session, user: User, matter_id: int, payload: bytes, extension: str, content_type: str, req: DemandLetterRequest, preview: DemandLetterPreview, visibility: str) -> MatterDocument:
    stored_name = f"{secrets.token_hex(24)}.{extension}"
    (vault_path / stored_name).write_bytes(payload)
    original_name = f"{_safe_stem(req)}.{extension}"
    item = MatterDocument(
        matter_id=matter_id,
        category="demand",
        title=f"Formal Demand - {req.verification.debtor_name}",
        description=f"Generated by Chambers Document Studio. Verification hash: {preview.verification.snapshot.inputs_hash}",
        original_name=original_name,
        stored_name=stored_name,
        content_type=content_type,
        size_bytes=len(payload),
        checksum_sha256=hashlib.sha256(payload).hexdigest(),
        visibility=visibility,
        uploaded_by_id=user.id,
    )
    db.add(item)
    db.flush()
    return item


@router.post("/demand-letter/preview", response_model=DemandLetterPreview)
def preview_demand_letter(req: DemandLetterRequest, _user: User = Depends(current_user)) -> DemandLetterPreview:
    return build_demand_letter(req)


@router.post("/demand-letter/render-html")
def render_demand_letter_html(req: DemandLetterRequest, _user: User = Depends(current_user)):
    result = build_demand_letter(req)
    if not result.eligible or not result.html:
        raise HTTPException(status_code=409, detail={"message": "Demand letter is blocked by debt verification", "blockers": result.blockers})
    return {"html": result.html, "verification_hash": result.verification.snapshot.inputs_hash, "deadline": result.deadline.isoformat()}


@router.post("/demand-letter/render-docx")
def render_demand_letter_docx(req: DemandLetterRequest, _user: User = Depends(current_user)):
    payload, _preview = build_docx_bytes(req)
    return StreamingResponse(io.BytesIO(payload), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="{_safe_stem(req)}.docx"'})


@router.post("/demand-letter/render-pdf")
def render_demand_letter_pdf(req: DemandLetterRequest, _user: User = Depends(current_user)):
    payload, _preview = build_pdf_bytes(req)
    return StreamingResponse(io.BytesIO(payload), media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{_safe_stem(req)}.pdf"'})


@router.post("/demand-letter/save-to-matter", status_code=201)
def save_demand_letter_to_matter(payload: SaveDemandLetterRequest, request: Request, db: Session = Depends(get_db), user: User = Depends(require_recovery("document:create"))):
    from app.operations import Matter

    matter = db.get(Matter, payload.matter_id)
    if not matter:
        raise HTTPException(status_code=404, detail="Matter not found")

    req = DemandLetterRequest(**payload.model_dump(exclude={"matter_id", "visibility", "formats"}))
    formats = list(dict.fromkeys(payload.formats))
    if not formats:
        raise HTTPException(status_code=422, detail="At least one output format is required")

    docx_bytes, preview = build_docx_bytes(req)
    created: list[MatterDocument] = []
    if "docx" in formats:
        created.append(_store_generated(db=db, user=user, matter_id=payload.matter_id, payload=docx_bytes, extension="docx", content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", req=req, preview=preview, visibility=payload.visibility))
    if "pdf" in formats:
        pdf_bytes, _ = build_pdf_bytes(req)
        created.append(_store_generated(db=db, user=user, matter_id=payload.matter_id, payload=pdf_bytes, extension="pdf", content_type="application/pdf", req=req, preview=preview, visibility=payload.visibility))

    audit(db, request, user, "document_studio.demand_saved", "legal_matter", str(payload.matter_id), after={"chambers_reference": req.chambers_reference, "debtor": req.verification.debtor_name, "source_reference": req.verification.source_reference, "verification_hash": preview.verification.snapshot.inputs_hash, "formats": formats, "document_ids": [item.id for item in created]})
    db.commit()
    return {
        "matter_id": payload.matter_id,
        "verification_hash": preview.verification.snapshot.inputs_hash,
        "documents": [{"id": item.id, "title": item.title, "original_name": item.original_name, "content_type": item.content_type, "checksum_sha256": item.checksum_sha256, "visibility": item.visibility, "download_url": f"/api/v1/ops/documents/{item.id}/download"} for item in created],
    }
