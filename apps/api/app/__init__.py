"""Lelefa Chambers API package bootstrap helpers."""

from __future__ import annotations

import zipfile
from pathlib import Path


def _ensure_demand_template() -> None:
    """Keep Document Studio usable even if a binary checkout is damaged.

    The repository carries the approved Batlokoa Word template. This fallback is
    intentionally modelled on the same supplied Lelefa Chambers letterhead and
    footer and is only generated when that package is missing or is not a valid
    OOXML/ZIP document.
    """
    template = Path(__file__).resolve().parent / "templates" / "batlokoa_formal_demand.docx"
    if template.exists() and zipfile.is_zipfile(template):
        return

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    template.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    navy = RGBColor(23, 59, 92)
    gold = "B78A34"

    header = section.header
    header.is_linked_to_previous = False
    title = header.paragraphs[0]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("LELEFA CHAMBERS")
    run.bold = True
    run.font.name = "Georgia"
    run.font.size = Pt(24)
    run.font.color.rgb = navy
    subtitle = title.add_run("\nAdvocates & Legal Practitioners")
    subtitle.italic = True
    subtitle.font.name = "Georgia"
    subtitle.font.size = Pt(10)
    subtitle.font.color.rgb = RGBColor(183, 138, 52)

    contact = header.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_run = contact.add_run(
        "Lenyora House, Office No. 4, 190 Nightingale Road, New Europa, Maseru, Lesotho | "
        "Matsepelelefa15@gmail.com | +266 5776 3829"
    )
    contact_run.font.name = "Georgia"
    contact_run.font.size = Pt(8)
    contact_run.font.color.rgb = navy
    ppr = contact._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), gold)
    borders.append(bottom)
    ppr.append(borders)

    body_lines = [
        "TO:\nNTHAKO LIMAKATSO\nNational ID: ____________________\nSource Ref: B1406",
        "27 September 2026",
        "Our Ref:\nLC/BATL/B1406/270926",
        "RE: FORMAL DEMAND FOR PAYMENT OF OUTSTANDING LOAN - M __________",
        "Dear Sir/Madam,",
        "(1) We act on the instructions of Batlokoa, our Client, in relation to the loan account held in your name.",
        "(2) Our Client's records reflect that, on or about 05 July 2024, a loan facility in the principal amount of M 10,000.00 was advanced to you.",
        "(3) The facility attracted interest at the agreed contractual rate of 20% per annum. Payments totalling M 1,600.00 have been credited to the account. Following application of those payments and contractual interest accrued under the loan agreement, the amount presently due and owing is M __________.",
        "(4) Despite prior recovery efforts undertaken on behalf of our Client, the above balance remains unpaid.",
        "(5) We therefore demand that, within seven (7) days of receipt of this letter, you either settle the outstanding balance of M __________ in full or contact Lelefa Chambers in writing to propose a payment arrangement acceptable to our Client.",
        "(6) If payment or a satisfactory written arrangement is not made within that period, we are instructed to advise our Client on further recovery action, which may include appropriate legal proceedings for recovery of the debt, together with contractual interest and recoverable legal costs, to the extent lawfully claimable.",
        "(7) If you dispute the balance or contend that any payment has not been credited, please provide documentary proof to our Chambers within the same seven (7)-day period. This letter constitutes a formal demand and is issued without waiver of any rights or remedies available to our Client.",
        "ACCOUNT SUMMARY",
    ]
    for text in body_lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.add_run(text)

    table = document.add_table(rows=3, cols=4)
    values = [
        ("Loan date", "05 July 2024", "Original principal", "M 10,000.00"),
        ("Interest rate", "20% per annum (contractual)", "Payments credited", "M 1,600.00"),
        ("Outstanding", "M __________", "Demand period", "7 days from receipt"),
    ]
    for row, values_row in zip(table.rows, values):
        for index, value in enumerate(values_row):
            cell = row.cells[index]
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Georgia"
                    run.font.size = Pt(9)
                    if index in (0, 2):
                        run.bold = True
                        run.font.color.rgb = navy

    document.add_paragraph("\nYours faithfully,")
    sign = document.add_paragraph()
    sign.add_run("_______________________________\n").bold = True
    sig = sign.add_run("Advocate Mats'epe Lelefa, LLM")
    sig.bold = True
    sig.font.color.rgb = navy

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run(
        "LELEFA CHAMBERS | Lenyora House, Office No. 4, 190 Nightingale Road, New Europa, Maseru | +266 5776 3829"
    )
    fr.font.name = "Georgia"
    fr.font.size = Pt(7)
    fr.font.color.rgb = navy
    fppr = fp._p.get_or_add_pPr()
    fborders = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "10")
    top.set(qn("w:space"), "3")
    top.set(qn("w:color"), gold)
    fborders.append(top)
    fppr.append(fborders)

    document.save(template)


_ensure_demand_template()
